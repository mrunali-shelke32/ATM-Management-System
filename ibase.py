
import openpyxl
import os

file_name = os.path.join(os.path.dirname(os.path.abspath(__file__)), "atm.xlsx")

if os.path.exists(file_name):
    wb = openpyxl.load_workbook(file_name)
else:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Phone Number", "Name", "Balance"])
    wb.save(file_name)

sheet = wb.active

while True:
    def menu():
        print("\n===== ATM MANAGEMENT SYSTEM =====")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")

    menu()
    choice = int(input("Enter your choice: "))

    if choice == 1:
        def create_account():
            phone = input("Enter your phone number: ")
            name = input("Enter your name: ")
            balance = int(input("Enter initial deposit: "))

            if balance < 0:
                print("Invalid deposit amount")
                return

            sheet.append([phone, name, balance])
            wb.save("atm.xlsx")
            print("Account created successfully!")

        create_account()

    elif choice == 2:
        def deposit():
            phone = input("Enter your phone number: ")
            amount = int(input("Enter deposit amount: "))

            if amount <= 0:
                print("Please enter a valid amount")
                return

            for row in sheet.iter_rows(min_row=2):
                if str(row[0].value) == phone:
                    row[2].value += amount
                    wb.save("atm.xlsx")
                    print("Deposit completed successfully!")
                    print("Updated balance:", row[2].value)
                    return

            print("Account not found!")

        deposit()

    elif choice == 3:
        def withdraw():
            phone = input("Enter your phone number: ")
            amount = int(input("Enter withdrawal amount: "))

            for row in sheet.iter_rows(min_row=2):
                if str(row[0].value) == phone:
                    balance = row[2].value

                    if amount <= 0:
                        print("Please enter a valid amount")
                    elif amount > balance:
                        print("Insufficient balance!")
                    else:
                        row[2].value = balance - amount
                        wb.save("atm.xlsx")
                        print("Withdrawal completed!")
                        print("Available balance:", row[2].value)

                    return

            print("Account not found!")

        withdraw()

    elif choice == 4:
        def check_balance():
            phone = input("Enter your phone number: ")

            for row in sheet.iter_rows(min_row=2, values_only=True):
                if str(row[0]) == phone:
                    print("Current balance:", row[2])
                    return

            print("Account not found!")

        check_balance()

    elif choice == 5:
        print("Thank you for using our ATM!")
        break

    else:
        print("Invalid choice. Please select 1 to 5.")
